# -*- coding: utf-8 -*-
"""
物料价格评审工具 v3.8
====================
基于 PySide6 的桌面应用，用于维护物料标准价格库、加载估算书、
自动比对价格差异并导出评审报告。

v3.8 优化:
- 加载估算书时，即使物料编码为空也照常读取（仅过滤合计/小计行与完全空行）
- 支持估算书"甲供"页签没有"物料编码"列的情况（仅需"名称"+"单价"两列）

v3.7 优化:
- 评审时若估算书物料缺少编码，自动改用物料名称进行模糊匹配
  （精确匹配 → 包含匹配 → 相似度匹配）

v3.6 优化:
- 将"物料名称"列设置为 Stretch 模式，自动拉伸填满剩余空间
- 其他列根据内容自适应宽度，且支持手动拖动调整
"""

import sys
import sqlite3
import os
import difflib
from datetime import datetime

from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QTabWidget, QVBoxLayout, QHBoxLayout,
    QLineEdit, QPushButton, QTableWidget, QTableWidgetItem, QLabel,
    QDialog, QFormLayout, QDialogButtonBox, QMessageBox, QFileDialog,
    QHeaderView, QCheckBox, QProgressDialog
)
from PySide6.QtCore import Qt, QCoreApplication
from PySide6.QtGui import QColor
import pandas as pd

# ======================== 数据库辅助类 ========================
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(SCRIPT_DIR, "price_library.db")


def _normalize_material_name(name):
    """规范化物料名称用于模糊匹配：去空白、统一乘号等写法、转小写。"""
    if name is None:
        return ''
    s = str(name).strip().lower()
    s = ''.join(s.split())
    for ch in ('×', '✕', '＊', '﹡', 'Ｘ', 'ｘ'):
        s = s.replace(ch, 'x')
    return s


class DBHelper:
    @staticmethod
    def init_db():
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS price_library (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                material_code TEXT UNIQUE NOT NULL,
                material_name TEXT NOT NULL,
                unit_price REAL NOT NULL,
                effective_date TEXT,
                status TEXT DEFAULT '有效'
            )
        ''')
        cursor.execute("SELECT COUNT(*) FROM price_library")
        if cursor.fetchone()[0] == 0:
            samples = [
                ('A001', '螺丝 M4x10', 0.85, '2026-01-01', '有效'),
                ('A002', 'PCB 主板 V3.0', 125.00, '2026-01-01', '有效'),
                ('B001', '电源线 1.5m', 12.50, '2026-01-01', '有效'),
                ('C001', '铝合金外壳', 45.00, '2026-01-01', '有效'),
            ]
            cursor.executemany(
                "INSERT INTO price_library (material_code, material_name, unit_price, effective_date, status) "
                "VALUES (?,?,?,?,?)",
                samples
            )
        conn.commit()
        conn.close()

    @staticmethod
    def get_connection():
        return sqlite3.connect(DB_PATH)

    @staticmethod
    def query_prices(keyword=''):
        conn = DBHelper.get_connection()
        cursor = conn.cursor()
        if keyword:
            sql = ("SELECT id, material_code, material_name, unit_price, effective_date "
                   "FROM price_library WHERE material_code LIKE ? OR material_name LIKE ? ORDER BY id")
            params = (f'%{keyword}%', f'%{keyword}%')
        else:
            sql = ("SELECT id, material_code, material_name, unit_price, effective_date "
                   "FROM price_library ORDER BY id")
            params = ()
        cursor.execute(sql, params)
        rows = cursor.fetchall()
        conn.close()
        return rows

    @staticmethod
    def update_price(record_id, field, value):
        allowed_fields = {"material_code", "material_name", "unit_price", "effective_date"}
        if field not in allowed_fields:
            raise ValueError(f"不允许更新的字段: {field}")
        conn = DBHelper.get_connection()
        cursor = conn.cursor()
        cursor.execute(f"UPDATE price_library SET {field}=? WHERE id=?", (value, record_id))
        conn.commit()
        conn.close()

    @staticmethod
    def delete_prices(record_ids):
        if not record_ids:
            return
        conn = DBHelper.get_connection()
        cursor = conn.cursor()
        placeholders = ','.join('?' * len(record_ids))
        cursor.execute(f"DELETE FROM price_library WHERE id IN ({placeholders})", record_ids)
        conn.commit()
        conn.close()

    @staticmethod
    def delete_price(record_id):
        conn = DBHelper.get_connection()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM price_library WHERE id=?", (record_id,))
        conn.commit()
        conn.close()

    @staticmethod
    def insert_price(code, name, price, date):
        conn = DBHelper.get_connection()
        cursor = conn.cursor()
        cursor.execute(
            "INSERT INTO price_library (material_code, material_name, unit_price, effective_date) "
            "VALUES (?,?,?,?)",
            (code, name, price, date)
        )
        conn.commit()
        conn.close()

    @staticmethod
    def get_latest_price(code):
        conn = DBHelper.get_connection()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT material_name, unit_price FROM price_library "
            "WHERE material_code=? AND status='有效' ORDER BY effective_date DESC LIMIT 1",
            (code,)
        )
        row = cursor.fetchone()
        conn.close()
        return row

    @staticmethod
    def get_all_effective_prices():
        """一次性读取价格库中所有有效物料，供名称模糊匹配复用。"""
        conn = DBHelper.get_connection()
        cursor = conn.cursor()
        cursor.execute(
            "SELECT material_code, material_name, unit_price FROM price_library "
            "WHERE status='有效' ORDER BY effective_date DESC"
        )
        rows = cursor.fetchall()
        conn.close()
        return rows

    @staticmethod
    def fuzzy_match_price(name, library=None, threshold=0.6):
        """
        根据物料名称在价格库中进行模糊匹配（用于估算书缺少物料编码的场景）。
        匹配优先级：精确匹配 → 包含匹配 → 相似度匹配(difflib)。
        返回 (material_code, material_name, unit_price)，未匹配返回 None。
        """
        target = _normalize_material_name(name)
        if not target:
            return None
        if library is None:
            library = DBHelper.get_all_effective_prices()

        # 1. 精确匹配
        for code, db_name, db_price in library:
            if _normalize_material_name(db_name) == target:
                return (code, db_name, db_price)

        # 2. 包含关系匹配
        for code, db_name, db_price in library:
            norm = _normalize_material_name(db_name)
            if norm and (target in norm or norm in target):
                return (code, db_name, db_price)

        # 3. 相似度匹配
        best = None
        best_score = 0.0
        for code, db_name, db_price in library:
            norm = _normalize_material_name(db_name)
            if not norm:
                continue
            score = difflib.SequenceMatcher(None, target, norm).ratio()
            if score > best_score:
                best_score = score
                best = (code, db_name, db_price)
        if best is not None and best_score >= threshold:
            return best
        return None


# ======================== 主窗口 ========================
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("物料价格评审工具 v3.8")
        self.resize(1200, 750)

        self.tabs = QTabWidget()
        self.setCentralWidget(self.tabs)

        self.maintain_tab = QWidget()
        self.tabs.addTab(self.maintain_tab, "📦 价格库维护")
        self.setup_maintain_tab()

        self.estimate_tab = QWidget()
        self.tabs.addTab(self.estimate_tab, "📄 估算书管理")
        self.setup_estimate_tab()

        self.review_tab = QWidget()
        self.tabs.addTab(self.review_tab, "⚖️ 价格评审")
        self.setup_review_tab()

        self.estimate_df = None
        self.review_full_df = None

        self.load_price_data()

    # ======================== 工具方法：自适应列宽 ========================
    def auto_resize_columns(self, table_widget, name_col_index=2):
        """
        自适应表格列宽：
        - "物料名称"列设置为 Stretch 模式，自动拉伸填满剩余空间
        - 其他列根据内容自适应宽度，且支持手动拖动调整
        
        Args:
            table_widget: 要调整的表格控件
            name_col_index: "物料名称"列的索引（默认2，即第3列）
        """
        header = table_widget.horizontalHeader()
        
        # 先根据内容调整所有列宽
        table_widget.resizeColumnsToContents()
        
        # 设置最小宽度
        for col in range(table_widget.columnCount()):
            current_width = table_widget.columnWidth(col)
            if current_width < 60:
                table_widget.setColumnWidth(col, 60)
        
        # 设置列宽模式
        for col in range(table_widget.columnCount()):
            if col == name_col_index:
                # "物料名称"列设置为 Stretch，自动填满剩余空间
                header.setSectionResizeMode(col, QHeaderView.ResizeMode.Stretch)
            else:
                # 其他列设置为 Interactive，用户可手动拖动调整
                header.setSectionResizeMode(col, QHeaderView.ResizeMode.Interactive)

    # ======================== 价格库维护 ========================
    def setup_maintain_tab(self):
        layout = QVBoxLayout(self.maintain_tab)

        top_bar = QHBoxLayout()
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("搜索物料编码/名称...")
        self.search_input.textChanged.connect(self.load_price_data)
        top_bar.addWidget(self.search_input)

        btn_add = QPushButton("➕ 新增")
        btn_add.clicked.connect(self.add_price)
        top_bar.addWidget(btn_add)

        btn_delete = QPushButton("🗑️ 删除选中")
        btn_delete.clicked.connect(self.delete_prices)
        top_bar.addWidget(btn_delete)

        btn_select_all = QPushButton("☑️ 全选")
        btn_select_all.clicked.connect(self.select_all_rows)
        top_bar.addWidget(btn_select_all)

        btn_deselect_all = QPushButton("☐ 取消全选")
        btn_deselect_all.clicked.connect(self.deselect_all_rows)
        top_bar.addWidget(btn_deselect_all)

        btn_refresh = QPushButton("🔄 刷新")
        btn_refresh.clicked.connect(self.load_price_data)
        top_bar.addWidget(btn_refresh)

        btn_import = QPushButton("📥 导入Excel标准价")
        btn_import.clicked.connect(self.import_excel)
        top_bar.addWidget(btn_import)
        top_bar.addStretch()
        layout.addLayout(top_bar)

        self.table = QTableWidget()
        self.table.setColumnCount(6)
        self.table.setHorizontalHeaderLabels(["", "ID", "物料编码", "物料名称", "单价", "生效日期"])
        # 初始设置为 Interactive，让用户可以拖动调整列宽
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.table.setEditTriggers(QTableWidget.EditTrigger.DoubleClicked)
        self.table.itemChanged.connect(self.on_cell_changed)
        self.checkbox_states = {}
        layout.addWidget(self.table)

        self.status_label = QLabel("就绪")
        layout.addWidget(self.status_label)

    def load_price_data(self):
        self.table.itemChanged.disconnect(self.on_cell_changed)
        rows = DBHelper.query_prices(self.search_input.text().strip())
        self.table.setRowCount(len(rows))
        self.checkbox_states.clear()

        for row_idx, row_data in enumerate(rows):
            checkbox = QCheckBox()
            checkbox.setChecked(False)
            checkbox.setProperty("row_id", row_data[0])
            checkbox.stateChanged.connect(lambda state, r=row_idx: self.on_checkbox_changed(r, state))
            self.table.setCellWidget(row_idx, 0, checkbox)
            self.checkbox_states[row_idx] = False

            for col_idx, val in enumerate(row_data):
                item = QTableWidgetItem(str(val))
                if col_idx == 0:
                    item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                self.table.setItem(row_idx, col_idx + 1, item)
        self.status_label.setText(f"共 {len(rows)} 条记录")
        self.table.itemChanged.connect(self.on_cell_changed)

        # 自适应列宽："物料名称"列（索引3）设置为 Stretch
        self.auto_resize_columns(self.table, name_col_index=3)

    def on_checkbox_changed(self, row, state):
        self.checkbox_states[row] = (state == Qt.CheckState.Checked.value)

    def select_all_rows(self):
        for row in range(self.table.rowCount()):
            checkbox = self.table.cellWidget(row, 0)
            if checkbox is not None:
                checkbox.setChecked(True)
                self.checkbox_states[row] = True

    def deselect_all_rows(self):
        for row in range(self.table.rowCount()):
            checkbox = self.table.cellWidget(row, 0)
            if checkbox is not None:
                checkbox.setChecked(False)
                self.checkbox_states[row] = False

    def get_selected_ids(self):
        selected_ids = []
        for row in range(self.table.rowCount()):
            if self.checkbox_states.get(row, False):
                id_item = self.table.item(row, 1)
                if id_item is not None:
                    try:
                        selected_ids.append(int(id_item.text()))
                    except ValueError:
                        pass
        return selected_ids

    def on_cell_changed(self, item):
        row, col = item.row(), item.column()
        if col == 0:
            return
        record_id_item = self.table.item(row, 1)
        if record_id_item is None:
            return
        try:
            record_id = int(record_id_item.text())
        except ValueError:
            return
        cols = ["", "id", "material_code", "material_name", "unit_price", "effective_date"]
        field = cols[col]
        new_value = item.text()
        try:
            if field == "unit_price":
                new_value = float(new_value)
            DBHelper.update_price(record_id, field, new_value)
            self.status_label.setText(f"✅ 更新成功: ID={record_id}, {field}={new_value}")
        except Exception as e:
            self.status_label.setText(f"❌ 更新失败: {e}")
            self.load_price_data()

    def add_price(self):
        dialog = QDialog(self)
        dialog.setWindowTitle("新增物料价格")
        layout = QFormLayout(dialog)
        code_edit = QLineEdit()
        name_edit = QLineEdit()
        price_edit = QLineEdit()
        date_edit = QLineEdit(datetime.now().strftime("%Y-%m-%d"))
        layout.addRow("物料编码*:", code_edit)
        layout.addRow("物料名称*:", name_edit)
        layout.addRow("单价*:", price_edit)
        layout.addRow("生效日期:", date_edit)
        btn_box = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        btn_box.accepted.connect(dialog.accept)
        btn_box.rejected.connect(dialog.reject)
        layout.addRow(btn_box)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            try:
                code = code_edit.text().strip()
                name = name_edit.text().strip()
                price = float(price_edit.text().strip())
                date = date_edit.text().strip() or datetime.now().strftime("%Y-%m-%d")
                if not code or not name:
                    raise ValueError("编码和名称不能为空")
                DBHelper.insert_price(code, name, price, date)
                self.load_price_data()
                self.status_label.setText(f"✅ 新增成功: {code}")
            except Exception as e:
                QMessageBox.critical(self, "错误", f"新增失败:\n{e}")

    def delete_prices(self):
        selected_ids = self.get_selected_ids()
        if not selected_ids:
            QMessageBox.warning(self, "提示", "请先勾选要删除的行（点击行首复选框）")
            return
        reply = QMessageBox.question(
            self, "确认删除",
            f"确定要删除选中的 {len(selected_ids)} 条记录吗？\n此操作不可撤销！"
        )
        if reply == QMessageBox.StandardButton.Yes:
            try:
                DBHelper.delete_prices(selected_ids)
                self.load_price_data()
                self.status_label.setText(f"🗑️ 已删除 {len(selected_ids)} 条记录")
            except Exception as e:
                QMessageBox.critical(self, "错误", f"删除失败:\n{e}")

    def import_excel(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, "选择Excel价格标准表", "", "Excel文件 (*.xlsx *.xls)"
        )
        if not file_path:
            return
        try:
            df = pd.read_excel(file_path)
            col_map = {}
            for col in df.columns:
                col_str = str(col)
                col_lower = col_str.lower()
                if '编码' in col_str or 'code' in col_lower:
                    col_map['code'] = col
                elif '名称' in col_str or 'name' in col_lower:
                    col_map['name'] = col
                elif '单价' in col_str or '价格' in col_str or 'price' in col_lower:
                    col_map['price'] = col
                elif '生效日期' in col_str or '日期' in col_str or 'date' in col_lower:
                    col_map['date'] = col

            if not all(k in col_map for k in ('code', 'name', 'price')):
                raise ValueError("Excel必须包含'物料编码/名称/单价'三列")

            conn = DBHelper.get_connection()
            cursor = conn.cursor()
            count = 0
            default_date = datetime.now().strftime("%Y-%m-%d")

            for _, row in df.iterrows():
                raw_code = row[col_map['code']]
                if isinstance(raw_code, float) and raw_code == int(raw_code):
                    code = str(int(raw_code)).strip()
                else:
                    code = str(raw_code).strip()
                name = str(row[col_map['name']]).strip()
                price = float(row[col_map['price']])

                if 'date' in col_map:
                    raw_date = row[col_map['date']]
                    if pd.isna(raw_date):
                        date = default_date
                    elif isinstance(raw_date, datetime):
                        date = raw_date.strftime("%Y-%m-%d")
                    else:
                        date = str(raw_date).strip()
                else:
                    date = default_date

                cursor.execute("SELECT id FROM price_library WHERE material_code=?", (code,))
                if cursor.fetchone():
                    cursor.execute(
                        "UPDATE price_library SET material_name=?, unit_price=?, effective_date=? WHERE material_code=?",
                        (name, price, date, code)
                    )
                else:
                    cursor.execute(
                        "INSERT INTO price_library (material_code, material_name, unit_price, effective_date) "
                        "VALUES (?,?,?,?)",
                        (code, name, price, date)
                    )
                count += 1

            conn.commit()
            conn.close()
            self.load_price_data()
            self.status_label.setText(f"✅ 导入/更新成功，处理 {count} 条记录")
        except Exception as e:
            QMessageBox.critical(self, "导入失败", str(e))

    # ======================== 估算书管理 ========================
    def setup_estimate_tab(self):
        layout = QVBoxLayout(self.estimate_tab)

        bar = QHBoxLayout()
        btn_load = QPushButton("📂 加载估算书Excel")
        btn_load.clicked.connect(self.load_estimate)
        bar.addWidget(btn_load)

        btn_clear = QPushButton("🗑️ 清空")
        btn_clear.clicked.connect(self.clear_estimate)
        bar.addWidget(btn_clear)

        self.estimate_search = QLineEdit()
        self.estimate_search.setPlaceholderText("在估算书中搜索...")
        self.estimate_search.textChanged.connect(self.filter_estimate)
        bar.addWidget(self.estimate_search)
        bar.addStretch()
        layout.addLayout(bar)

        self.estimate_table = QTableWidget()
        self.estimate_table.setColumnCount(0)
        self.estimate_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        layout.addWidget(self.estimate_table)

        self.estimate_status = QLabel("未加载估算书")
        layout.addWidget(self.estimate_status)

    def load_estimate(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, "选择估算书Excel", "", "Excel文件 (*.xlsx *.xls)"
        )
        if not file_path:
            return
        try:
            self.estimate_df = self._load_jigong_estimate(file_path)
            src = getattr(self, '_estimate_source', None)
            msg = f"✅ 已加载: {os.path.basename(file_path)}，共 {len(self.estimate_df)} 行"
            if src:
                msg += f"（来源页签: {src}）"
            self.estimate_status.setText(msg)
            self.display_estimate(self.estimate_df)
        except Exception as e:
            QMessageBox.critical(self, "错误", f"加载失败: {e}")

    def _load_jigong_estimate(self, path):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        jigong_sheets = [s for s in wb.sheetnames if '甲供' in s]
        if not jigong_sheets:
            wb.close()
            df = pd.read_excel(path)
            if df is None or len(df) == 0:
                raise ValueError("未找到'甲供'页签，且文件无可读数据")
            self._estimate_source = None
            return df
        records = []
        self._estimate_source = '、'.join(jigong_sheets)
        for sname in jigong_sheets:
            ws = wb[sname]
            header_row = None
            # 优先用"物料编码"定位表头行（最独特，避免误判标题行）
            for r in range(1, min(ws.max_row, 30) + 1):
                for c in range(1, ws.max_column + 1):
                    v = ws.cell(r, c).value
                    if v and '物料编码' in str(v):
                        header_row = r
                        break
                if header_row:
                    break
            # 若无"物料编码"列，退而用名称/单价列定位表头行
            if not header_row:
                for r in range(1, min(ws.max_row, 30) + 1):
                    for c in range(1, ws.max_column + 1):
                        v = ws.cell(r, c).value
                        sv = str(v) if v is not None else ''
                        if ('材料名称' in sv) or ('物料名称' in sv) or (sv == '单价'):
                            header_row = r
                            break
                    if header_row:
                        break
            if not header_row:
                continue
            col_code = col_name = col_price = None
            for c in range(1, ws.max_column + 1):
                v = ws.cell(header_row, c).value
                v2 = ws.cell(header_row + 1, c).value if header_row + 1 <= ws.max_row else None
                sv = str(v) if v is not None else ''
                sv2 = str(v2) if v2 is not None else ''
                if '物料编码' in sv or sv == '编码':
                    col_code = c
                elif '材料名称' in sv or '物料名称' in sv:
                    col_name = c
                elif '单价' in sv or '单价' in sv2:
                    col_price = c
            # 必须至少有"名称"与"单价"两列；"编码"列可选
            if col_name is None or col_price is None:
                continue
            for r in range(header_row + 2, ws.max_row + 1):
                code = ws.cell(r, col_code).value if col_code is not None else None
                name = ws.cell(r, col_name).value
                price = ws.cell(r, col_price).value

                code_str = '' if code is None else str(code).strip()
                name_str = '' if name is None else str(name).strip()

                # 过滤合计/小计行
                if name_str and ('合计' in name_str or '小计' in name_str):
                    continue

                # 过滤完全空行（物料编码与名称均为空）
                if not code_str and not name_str:
                    continue

                try:
                    price_f = float(price) if price is not None else 0.0
                except (ValueError, TypeError):
                    price_f = 0.0

                records.append({
                    '物料编码': code_str,
                    '材料名称': name_str,
                    '单价': price_f,
                    '来源页签': sname,
                })
        wb.close()
        if not records:
            raise ValueError("在'甲供'页签中未找到有效数据")
        return pd.DataFrame(records)

    def display_estimate(self, df):
        self.estimate_table.setRowCount(len(df))
        self.estimate_table.setColumnCount(len(df.columns))
        self.estimate_table.setHorizontalHeaderLabels([str(c) for c in df.columns])
        for row_idx, (_, row) in enumerate(df.iterrows()):
            for col_idx, val in enumerate(row):
                item = QTableWidgetItem(str(val))
                self.estimate_table.setItem(row_idx, col_idx, item)
        # 自适应列宽：估算书表格的"材料名称"列（索引1）设置为 Stretch
        # 注意：估算书表格列顺序为 [物料编码, 材料名称, 单价, 来源页签]
        self.auto_resize_columns(self.estimate_table, name_col_index=1)

    def filter_estimate(self):
        if self.estimate_df is None:
            return
        keyword = self.estimate_search.text().strip()
        if not keyword:
            self.display_estimate(self.estimate_df)
        else:
            mask = self.estimate_df.apply(
                lambda row: row.astype(str).str.contains(keyword, case=False, na=False).any(), axis=1
            )
            filtered = self.estimate_df[mask]
            self.display_estimate(filtered)
            self.estimate_status.setText(f"筛选结果: {len(filtered)} / {len(self.estimate_df)} 行")

    def clear_estimate(self):
        self.estimate_df = None
        self.estimate_table.setRowCount(0)
        self.estimate_table.setColumnCount(0)
        self.estimate_status.setText("已清空估算书")

    # ======================== 价格评审 ========================
    def setup_review_tab(self):
        layout = QVBoxLayout(self.review_tab)

        bar = QHBoxLayout()
        btn_review = QPushButton("🚀 执行评审")
        btn_review.clicked.connect(self.run_review)
        bar.addWidget(btn_review)

        rule_label = QLabel("📋 规则: |差异%|≤5%合理 | 5%~10%预警 | >10%不合理")
        rule_label.setStyleSheet("color: #555; font-weight: bold;")
        bar.addWidget(rule_label)

        self.filter_combo = QCheckBox("仅显示异常项（预警/不合理/缺失）")
        self.filter_combo.stateChanged.connect(self.apply_review_filter)
        bar.addWidget(self.filter_combo)

        btn_export = QPushButton("📤 导出评审结果")
        btn_export.clicked.connect(self.export_review)
        bar.addWidget(btn_export)
        bar.addStretch()
        layout.addLayout(bar)

        self.review_table = QTableWidget()
        self.review_table.setColumnCount(8)
        self.review_table.setHorizontalHeaderLabels(
            ["物料编码", "物料名称", "估算书单价", "价格库单价", "差异额", "差异%", "评审结果", "状态说明"]
        )
        self.review_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        layout.addWidget(self.review_table)

        self.review_stats = QLabel("请先加载估算书，然后点击『执行评审』")
        self.review_stats.setStyleSheet("font-weight: bold; padding: 5px; background: #f0f0f0; border-radius: 3px;")
        layout.addWidget(self.review_stats)

        self.review_full_df = None

    def _classify_diff(self, diff_percent):
        """
        根据差异百分比分类评审结果
        返回: (评审结果, 状态说明, 颜色)
        颜色说明：
        - 合理: 绿色 (#d4edda)
        - 预警: 橙色 (#ffeaa7) 
        - 不合理: 红色 (#ffcccc)
        - 库中无此物料: 灰色 (#d3d3d3)
        """
        if diff_percent is None:
            return "库中无此物料", "价格库中未找到该物料", "#d3d3d3"
        abs_pct = abs(diff_percent)
        if abs_pct <= 5.0:
            return "合理", f"偏差{abs_pct:.2f}%≤5%，可接受", "#d4edda"
        elif abs_pct <= 10.0:
            return "预警", f"偏差{abs_pct:.2f}%在5%~10%之间，需关注", "#ffeaa7"
        else:
            return "不合理", f"偏差{abs_pct:.2f}%>10%，需调整", "#ffcccc"

    def run_review(self):
        if self.estimate_df is None:
            QMessageBox.warning(self, "提示", "请先在『估算书管理』中加载估算书Excel")
            return

        cols = self.estimate_df.columns.tolist()
        code_col, name_col, price_col = None, None, None
        for c in cols:
            c_str = str(c)
            c_lower = c_str.lower()
            if '编码' in c_str or 'code' in c_lower:
                code_col = c
            if '名称' in c_str or 'name' in c_lower:
                name_col = c
            if '单价' in c_str or '价格' in c_str or 'price' in c_lower:
                price_col = c
        if price_col is None:
            QMessageBox.critical(
                self, "错误",
                "无法自动识别'单价'列，请确保表头包含'单价/价格'相关字样"
            )
            return
        if code_col is None and name_col is None:
            QMessageBox.critical(
                self, "错误",
                "无法自动识别'物料编码'或'物料名称'列，请确保表头包含相关字样"
            )
            return

        total = len(self.estimate_df)
        progress = QProgressDialog("正在比对物料价格...", "取消", 0, total, self)
        progress.setWindowModality(Qt.WindowModality.WindowModal)
        progress.show()

        # 预加载价格库有效物料，供名称模糊匹配复用
        library = DBHelper.get_all_effective_prices()

        results = []
        for idx, (_, row) in enumerate(self.estimate_df.iterrows()):
            if progress.wasCanceled():
                break
            progress.setValue(idx + 1)
            QCoreApplication.processEvents()

            code = ''
            if code_col is not None:
                code_val = row[code_col]
                code = '' if pd.isna(code_val) else str(code_val).strip()

            name = ''
            if name_col is not None:
                name_val = row[name_col]
                name = '' if pd.isna(name_val) else str(name_val).strip()

            try:
                est_price = float(row[price_col])
            except (ValueError, TypeError):
                est_price = 0.0

            db_info = None
            db_code = ''
            db_name = ''
            db_price = 0.0
            matched_by_name = False

            if code:
                # 有物料编码：优先按编码精确查找
                info = DBHelper.get_latest_price(code)
                if info:
                    db_code, db_name, db_price = code, info[0], info[1]
                    db_info = info
            elif name:
                # 无物料编码：按物料名称模糊匹配
                info = DBHelper.fuzzy_match_price(name, library=library)
                if info:
                    db_code, db_name, db_price = info
                    db_info = info
                    matched_by_name = True

            if db_info:
                diff = est_price - db_price
                diff_percent = (diff / db_price * 100) if db_price != 0 else None
                result, desc, color = self._classify_diff(diff_percent)
                if matched_by_name:
                    desc = f"[名称匹配] {desc}"
                results.append([
                    db_code, db_name, round(est_price, 4), round(db_price, 4),
                    round(diff, 4) if diff is not None else "-",
                    f"{diff_percent:.2f}%" if diff_percent is not None else "-",
                    result, desc, color
                ])
            else:
                result, desc, color = self._classify_diff(None)
                if not code and name:
                    desc = f"[名称未匹配] {desc}"
                display_name = name if name else "未找到"
                results.append([
                    code, display_name, round(est_price, 4), "-",
                    "-", "-", result, desc, color
                ])

        progress.setValue(total)

        self.review_full_df = pd.DataFrame(
            results,
            columns=["物料编码", "物料名称", "估算书单价", "价格库单价", "差异额", "差异%", "评审结果", "状态说明", "颜色"]
        )
        self.apply_review_filter()

        total_items = len(results)
        stats = {"合理": 0, "预警": 0, "不合理": 0, "库中无此物料": 0}
        for r in results:
            stats[r[6]] += 1

        diff_sum = 0.0
        for r in results:
            if isinstance(r[4], (int, float)):
                diff_sum += abs(r[4])

        self.review_stats.setText(
            f"📊 总计: {total_items} 项 | 🟢 合理: {stats['合理']} | 🟡 预警: {stats['预警']} | "
            f"🔴 不合理: {stats['不合理']} | ⚪ 库中无此物料: {stats['库中无此物料']} | "
            f"差异总额(绝对值): {diff_sum:.2f}"
        )

        # 评审完成后自适应列宽：评审表格的"物料名称"列（索引1）设置为 Stretch
        self.auto_resize_columns(self.review_table, name_col_index=1)

    def apply_review_filter(self):
        if self.review_full_df is None or len(self.review_full_df) == 0:
            self.review_table.setRowCount(0)
            return

        if self.filter_combo.isChecked():
            filtered = self.review_full_df[
                self.review_full_df["评审结果"].isin(["预警", "不合理", "库中无此物料"])
            ]
        else:
            filtered = self.review_full_df
        self.display_review(filtered)

    def display_review(self, df):
        """显示评审结果，每行整行应用对应底色"""
        self.review_table.setRowCount(len(df))
        for row_idx, (_, row) in enumerate(df.iterrows()):
            color = row.iloc[8] if len(row) > 8 else "#ffffff"
            for col_idx in range(8):
                val = row.iloc[col_idx]
                item = QTableWidgetItem(str(val))
                item.setBackground(QColor(color))
                self.review_table.setItem(row_idx, col_idx, item)

        if len(df) == 0:
            self.review_table.setRowCount(1)
            for col in range(8):
                item = QTableWidgetItem("（无匹配记录）")
                item.setBackground(QColor("#f0f0f0"))
                self.review_table.setItem(0, col, item)

        # 自适应列宽：评审表格的"物料名称"列（索引1）设置为 Stretch
        self.auto_resize_columns(self.review_table, name_col_index=1)

    def export_review(self):
        if self.review_full_df is None or len(self.review_full_df) == 0:
            QMessageBox.warning(self, "提示", "没有评审结果可导出，请先执行评审")
            return
        file_path, _ = QFileDialog.getSaveFileName(
            self, "保存评审结果", "评审结果.xlsx", "Excel文件 (*.xlsx)"
        )
        if not file_path:
            return
        try:
            df_export = self.review_full_df.iloc[:, :8].copy()

            total_items = len(df_export)
            stats = df_export["评审结果"].value_counts().to_dict()
            summary_row = {
                "物料编码": "汇总",
                "物料名称": f"共{total_items}项",
                "估算书单价": f"合理:{stats.get('合理',0)}",
                "价格库单价": f"预警:{stats.get('预警',0)}",
                "差异额": f"不合理:{stats.get('不合理',0)}",
                "差异%": f"缺失:{stats.get('库中无此物料',0)}",
                "评审结果": "---",
                "状态说明": "---"
            }
            df_export = pd.concat([df_export, pd.DataFrame([summary_row])], ignore_index=True)
            df_export.to_excel(file_path, index=False)
            QMessageBox.information(self, "成功", f"结果已导出至:\n{file_path}")
        except Exception as e:
            QMessageBox.critical(self, "导出失败", str(e))


# ======================== 启动 ========================
if __name__ == "__main__":
    DBHelper.init_db()
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())